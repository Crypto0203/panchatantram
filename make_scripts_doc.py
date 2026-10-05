import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import json

def set_cell_background(cell, fill_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_script_doc():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    PRIMARY = RGBColor(0x1A, 0x36, 0x5D)
    SECONDARY = RGBColor(0x2B, 0x6C, 0xB0)
    ACCENT = RGBColor(0xC0, 0x56, 0x21)
    TEXT_DARK = RGBColor(0x2D, 0x37, 0x48)

    # Title
    p_title = doc.add_paragraph()
    r_title = p_title.add_run("🐰 PANCHATANTRA KIDS — MASTER COMPLETE SCRIPTS BOOK")
    r_title.bold = True
    r_title.font.size = Pt(22)
    r_title.font.color.rgb = PRIMARY

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("Full 30-Second Production Scripts Formatted into Exact 3 × 10-Second Clips (10s * 3)")
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = SECONDARY

    p_meta = doc.add_paragraph()
    r_meta = p_meta.add_run("Format: 3 Clips × 10 Seconds = 30.0s Total | Spoken Telugu Voiceover + English Subtitles + 3D Pixar Visual Prompts + SFX + Next-Episode Teasers")
    r_meta.font.size = Pt(9.5)
    r_meta.font.italic = True

    # Load episodes
    with open(r'C:\Users\Suresh\.gemini\antigravity-ide\scratch\panchatantra-dashboard\src\episodes.js', 'r', encoding='utf-8') as f:
        text = f.read()
        raw_json = text.replace('export const episodes = ', '').rstrip(';\n')
        eps = json.loads(raw_json)

    print(f"Adding complete scripts for {len(eps)} episodes...")

    for ep in eps:
        p_ep = doc.add_paragraph()
        p_ep.paragraph_format.space_before = Pt(16)
        p_ep.paragraph_format.space_after = Pt(2)
        p_ep.paragraph_format.keep_with_next = True
        
        r_ep = p_ep.add_run(f"EPISODE {ep['id']:02d}: {ep['title'].upper()}\n")
        r_ep.bold = True
        r_ep.font.size = Pt(14)
        r_ep.font.color.rgb = PRIMARY
        
        r_meta = p_ep.add_run(f"🐾 Characters: {ep['characters']} | 💡 Moral: {ep['moral']} | ⏱️ Total Runtime: 30.0s (3 × 10s)\n")
        r_meta.font.size = Pt(9.5)
        r_meta.font.bold = True
        r_meta.font.color.rgb = SECONDARY

        # Table for the 3 clips
        tbl = doc.add_table(rows=4, cols=5)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        widths = [0.9, 1.8, 1.6, 1.4, 1.3]
        headers = ["Clip (Time)", "AI Visual Action Prompt (9:16)", "🎙️ Telugu Spoken VO", "💬 English Subtitles", "🔊 Sound Foley & Music"]
        
        # Header row
        hdr = tbl.rows[0]
        for idx, (title, w) in enumerate(zip(headers, widths)):
            cell = hdr.cells[idx]
            cell.width = Inches(w)
            set_cell_background(cell, "1A365D")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(title)
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        # 3 Clips
        for c_idx, clip in enumerate(ep['clips']):
            row = tbl.rows[c_idx + 1]
            bg = "F7FAFC" if c_idx % 2 == 1 else "FFFFFF"
            
            c_vals = [
                f"CLIP {clip['clipNumber']}\n({clip['timeRange'].split()[0]}-{clip['timeRange'].split()[2]})\n\n{clip['purpose']}",
                clip['visualPrompt'],
                clip['teluguVO'],
                clip['englishSub'],
                clip['sfx']
            ]
            
            for col_idx, (val, w) in enumerate(zip(c_vals, widths)):
                cell = row.cells[col_idx]
                cell.width = Inches(w)
                set_cell_background(cell, bg)
                set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
                p = cell.paragraphs[0]
                r = p.add_run(val)
                r.font.size = Pt(8.5)
                r.font.color.rgb = TEXT_DARK
                if col_idx == 0:
                    r.bold = True
                elif col_idx == 2:
                    r.font.color.rgb = RGBColor(0x0C, 0x4A, 0x6E)
                elif col_idx == 3:
                    r.font.color.rgb = RGBColor(0x85, 0x4D, 0x0E)

        # Engagement line
        p_eng = doc.add_paragraph()
        p_eng.paragraph_format.space_before = Pt(4)
        p_eng.paragraph_format.space_after = Pt(10)
        r_eng = p_eng.add_run(f"💬 Comment Trigger: {ep['commentQ']} | 🔔 Next Episode Bridge: EP {ep['nextEpId']} ({ep['nextEpTitle']})")
        r_eng.font.size = Pt(9)
        r_eng.font.italic = True
        r_eng.font.color.rgb = ACCENT

    out_path = r'C:\Users\Suresh\Downloads\Panchatantra_Kids_Complete_30s_Production_Scripts.docx'
    doc.save(out_path)
    print("SUCCESS: Master Complete Scripts Book saved to:", out_path)

if __name__ == '__main__':
    generate_script_doc()
